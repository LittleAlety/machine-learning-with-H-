#!/usr/bin/env python3
"""Build the comparison paper as a formatted DOCX with embedded figures and table."""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(ROOT, "PAPER_v2_comparison.md")
FIG = os.path.join(ROOT, "figures_v2")
OUTDOC = os.path.join(ROOT, "PAPER_v2_comparison.docx")

TERRA = RGBColor(0x7c,0x35,0x1c); INK = RGBColor(0x2b,0x23,0x1c); GREY = RGBColor(0x5a,0x50,0x44)

doc = Document()

# Page setup
for sec in doc.sections:
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)

def force_font(style, family):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for attr in ("w:asciiTheme","w:hAnsiTheme","w:cstheme","w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii","w:hAnsi","w:cs","w:eastAsia"):
        rfonts.set(qn(attr), family)

# Base styles
normal = doc.styles["Normal"]
normal.font.size = Pt(11)
force_font(normal, "Times New Roman")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

def style_heading(name, size, color, bold=True, before=12, after=6):
    st = doc.styles[name]
    st.font.size = Pt(size); st.font.bold = bold
    st.font.color.rgb = color
    force_font(st, "Times New Roman")
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

style_heading("Heading 1", 14, TERRA)
style_heading("Heading 2", 12, TERRA)

def add_inline(par, text):
    """Render inline **bold**, *italic*, and strip $ math delimiters."""
    text = text.replace("$", "")
    # tokenize bold and italic
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        else:
            r = par.add_run(tok[1:-1]); r.italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])

def add_figure(fname, caption, width=5.9):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(FIG, fname), width=Inches(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = GREY
    c.paragraph_format.space_after = Pt(10)

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),hexcolor)
    tcPr.append(shd)

def add_table(rows):
    header, body = rows[0], rows[2:]  # rows[1] is the |---| separator
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for j,h in enumerate(header):
        hdr[j].text = ""
        par = hdr[j].paragraphs[0]; run = par.add_run(h.strip().replace("**",""))
        run.bold = True; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(0xff,0xff,0xff)
        set_cell_bg(hdr[j], "7c351c")
    for brow in body:
        cells = t.add_row().cells
        for j,val in enumerate(brow):
            if j >= len(cells): continue
            cells[j].text = ""
            par = cells[j].paragraphs[0]
            run = par.add_run(val.strip().replace("**","")); run.font.size = Pt(8.5)
            if j == 0: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Figures to insert after headings (keyword -> list)
FIG_PLAN = {
    "Main Model": [("fig1_parity.png","Figure 1. Out-of-fold parity of the H* main model over 1,836 surfaces."),
                   ("fig2_shap.png","Figure 2. Global SHAP feature attribution of the 84-descriptor v3 model."),
                   ("fig3_volcano.png","Figure 3. Predicted HER volcano; green band marks the optimal |ΔG|≤0.1 eV window.")],
    "External Validation": [("fig4_external.png","Figure 4. Zero-retraining external validation: CatHub homogeneous (left) and cross-database EqV2-HER (right).")],
    "Composition-Only Descriptors": [("fig5_saturation.png","Figure 5. Three-family static/bulk descriptor ablation; every change lies within ±0.0016 eV.")],
    "Preregistered Active-Learning": [("fig6_chgnet_artifact.png","Figure 6. Universal-potential artifact: CHGNet vs DFT top-layer displacement (left) and median ratio ≈40× (right).")],
}

with open(MD, encoding="utf-8") as f:
    lines = f.read().split("\n")

i = 0; first_h1 = True
while i < len(lines):
    line = lines[i].rstrip()
    # table block
    if line.startswith("|"):
        block = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            block.append(cells); i += 1
        add_table(block)
        continue
    if not line.strip():
        i += 1; continue
    if line.startswith("# "):
        # Title
        tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = tp.add_run(line[2:].strip()); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = TERRA
        tp.paragraph_format.space_after = Pt(14)
        i += 1; continue
    if line.startswith("## "):
        heading = line[3:].strip()
        doc.add_heading(heading, level=1)
        # insert figures after this heading (before its body)
        for kw, figs in FIG_PLAN.items():
            if kw in heading:
                for fn, cap in figs: add_figure(fn, cap)
        i += 1; continue
    if line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2); i += 1; continue
    # normal paragraph (may wrap multiple consecutive non-empty lines)
    buf = [line]; i += 1
    while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(("#","|")):
        buf.append(lines[i].rstrip()); i += 1
    text = " ".join(buf)
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(par, text)

doc.save(OUTDOC)
print("Saved:", OUTDOC)
