#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Chinese comparison paper as a formatted DOCX with embedded figures and table."""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(ROOT, "PAPER_v2_comparison_CN.md")
FIG = os.path.join(ROOT, "figures_v2")
OUTDOC = os.path.join(ROOT, "PAPER_v2_comparison_CN.docx")

TERRA = RGBColor(0x7c,0x35,0x1c); GREY = RGBColor(0x5a,0x50,0x44)

doc = Document()
for sec in doc.sections:
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)

def force_font(style, latin, cjk):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for attr in ("w:asciiTheme","w:hAnsiTheme","w:cstheme","w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn("w:ascii"), latin); rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin); rfonts.set(qn("w:eastAsia"), cjk)

normal = doc.styles["Normal"]
normal.font.size = Pt(11)
force_font(normal, "Times New Roman", "宋体")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3

def style_heading(name, size):
    st = doc.styles[name]
    st.font.size = Pt(size); st.font.bold = True
    st.font.color.rgb = TERRA
    force_font(st, "Times New Roman", "黑体")
    st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

style_heading("Heading 1", 14)
style_heading("Heading 2", 12)

def add_inline(par, text):
    text = text.replace("$", "")
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        else:
            r = par.add_run(tok[1:-1]); r.italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])

def add_figure(fname, caption, width=5.9):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(FIG, fname), width=Inches(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = GREY
    c.paragraph_format.space_after = Pt(10)

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),hexcolor)
    tcPr.append(shd)

def add_table(rows):
    header, body = rows[0], rows[2:]
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for j,h in enumerate(header):
        hdr[j].text = ""
        par = hdr[j].paragraphs[0]; run = par.add_run(h.strip().replace("**",""))
        run.bold = True; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(0xff,0xff,0xff)
        set_cell_bg(hdr[j], "7c351c")
    for brow in body:
        cells = t.add_row().cells
        for j,val in enumerate(brow):
            if j >= len(cells): continue
            cells[j].text = ""
            par = cells[j].paragraphs[0]
            run = par.add_run(val.strip().replace("**","")); run.font.size = Pt(8.5)
            if j == 0: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

FIG_PLAN = {
    "主线模型": [("fig1_parity.png","图 1. H* 主线模型在 1,836 个表面上的样本外奇偶对比。"),
               ("fig2_shap.png","图 2. 84 描述符 v3 模型的全局 SHAP 特征归因。"),
               ("fig3_volcano.png","图 3. 预测 HER 火山；绿色色带标记最优 |ΔG|≤0.1 eV 窗口。")],
    "外部验证": [("fig4_external.png","图 4. 零重训外部验证：Catalysis-Hub 同质金属（左）与跨数据库 EqV2-HER（右）。")],
    "纯成分描述符": [("fig5_saturation.png","图 5. 三家族静态/体相描述符消融；所有变化均在 ±0.0016 eV 内。")],
    "预注册主动学习": [("fig6_chgnet_artifact.png","图 6. 通用势伪影：CHGNet 与 DFT 顶层位移（左）及中位比约 40 倍（右）。")],
}

with open(MD, encoding="utf-8") as f:
    lines = f.read().split("\n")

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if line.startswith("|"):
        block = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            block.append(cells); i += 1
        add_table(block); continue
    if not line.strip():
        i += 1; continue
    if line.startswith("# "):
        tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = tp.add_run(line[2:].strip()); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = TERRA
        tp.paragraph_format.space_after = Pt(14); i += 1; continue
    if line.startswith("## "):
        heading = line[3:].strip()
        doc.add_heading(heading, level=1)
        for kw, figs in FIG_PLAN.items():
            if kw in heading:
                for fn, cap in figs: add_figure(fn, cap)
        i += 1; continue
    if line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2); i += 1; continue
    buf = [line]; i += 1
    while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(("#","|")):
        buf.append(lines[i].rstrip()); i += 1
    text = " ".join(buf)
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(par, text)

doc.save(OUTDOC)
print("Saved:", OUTDOC)
